from __future__ import annotations

import argparse
import copy
import json
import os
import re
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Mapping

from openpyxl import load_workbook


DEFAULT_TREE_PATH = Path(__file__).with_name("risk2file_tree.json")

SEQUENCE_PATTERN = re.compile(r"^(?P<category>[1-6])(?:\.\d+){3}$")

CATEGORY_CONFIG = {
    "1": {
        "sheet": "1.企业关联方合规风险指标",
        "risk_root": "企业关联方合规风险",
    },
    "2": {
        "sheet": "2.产品法律风险",
        "risk_root": "产品法律风险",
    },
    "3": {
        "sheet": "3.企业信用风险指标",
        "risk_root": "企业信用风险",
    },
    "4": {
        "sheet": "4.劳动法律风险",
        "risk_root": "劳动用工法律合规风险",
    },
    "5": {
        "sheet": "5.企业国际化经营法律风险",
        "risk_root": "企业国际化经营合规风险",
    },
    "6": {
        "sheet": "6.企业供应链法律风险",
        "risk_root": "供应链合规风险",
    },
}


@dataclass(frozen=True)
class WorkbookIndicator:
    sequence: str
    description: str
    sheet_name: str
    row_number: int
    risk_root: str


@dataclass
class TreeIndicator:
    sequence: str
    indicator_name: str
    risk_root: str
    leaf: dict[str, Any]


class IndicatorDescriptionSyncError(ValueError):
    def __init__(self, errors: list[str], report: Mapping[str, Any] | None = None) -> None:
        self.errors = errors
        self.report = dict(report or {})
        self.report["errors"] = errors
        super().__init__("; ".join(errors))


def _normalize_sheet_name(value: object) -> str:
    return re.sub(r"\s+", "", str(value or "").strip())


def _normalize_header(value: object) -> str:
    return re.sub(r"\s+", "", str(value or "").strip())


def _normalize_sequence(value: object) -> str:
    return str(value or "").strip()


def _normalize_description(value: object) -> str:
    return str(value or "").replace("\r\n", "\n").replace("\r", "\n").strip()


def _expected_sheet_index() -> dict[str, tuple[str, str]]:
    return {
        _normalize_sheet_name(config["sheet"]): (prefix, config["risk_root"])
        for prefix, config in CATEGORY_CONFIG.items()
    }


def _find_header_columns(row: tuple[Any, ...]) -> tuple[int, int] | None:
    sequence_column: int | None = None
    description_column: int | None = None
    for index, value in enumerate(row):
        header = _normalize_header(value)
        if header.startswith("指标编号"):
            sequence_column = index
        elif header.startswith("指标描述"):
            description_column = index
    if sequence_column is None or description_column is None:
        return None
    return sequence_column, description_column


def load_workbook_indicators(
    workbook_path: str | Path,
) -> tuple[dict[str, WorkbookIndicator], list[str], int]:
    path = Path(workbook_path)
    if not path.is_file():
        return {}, [f"指标工作簿不存在：{path}"], 0

    expected_sheets = _expected_sheet_index()
    workbook = load_workbook(path, read_only=True, data_only=True)
    errors: list[str] = []
    indicators: dict[str, WorkbookIndicator] = {}
    seen_sheets: set[str] = set()

    try:
        for worksheet in workbook.worksheets:
            normalized_sheet = _normalize_sheet_name(worksheet.title)
            if normalized_sheet not in expected_sheets:
                continue

            prefix, risk_root = expected_sheets[normalized_sheet]
            seen_sheets.add(normalized_sheet)
            header_columns: tuple[int, int] | None = None
            header_row = 0

            for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                if header_columns is None:
                    if row_number <= 10:
                        header_columns = _find_header_columns(row)
                        if header_columns is not None:
                            header_row = row_number
                    continue

                sequence_column, description_column = header_columns
                raw_sequence = row[sequence_column] if sequence_column < len(row) else None
                if raw_sequence is None or not str(raw_sequence).strip():
                    continue

                sequence = _normalize_sequence(raw_sequence)
                match = SEQUENCE_PATTERN.fullmatch(sequence)
                if match is None:
                    errors.append(
                        f"{worksheet.title}第{row_number}行的指标编号格式错误：{sequence}"
                    )
                    continue
                if match.group("category") != prefix:
                    errors.append(
                        f"{worksheet.title}第{row_number}行的指标编号{sequence}与工作表类别不一致"
                    )
                    continue

                raw_description = (
                    row[description_column] if description_column < len(row) else None
                )
                description = _normalize_description(raw_description)
                if not description:
                    errors.append(
                        f"{worksheet.title}第{row_number}行的指标{sequence}缺少指标描述"
                    )
                    continue
                if sequence in indicators:
                    previous = indicators[sequence]
                    errors.append(
                        f"指标编号{sequence}重复：{previous.sheet_name}第{previous.row_number}行、"
                        f"{worksheet.title}第{row_number}行"
                    )
                    continue

                indicators[sequence] = WorkbookIndicator(
                    sequence=sequence,
                    description=description,
                    sheet_name=worksheet.title,
                    row_number=row_number,
                    risk_root=risk_root,
                )

            if header_columns is None:
                errors.append(
                    f"工作表{worksheet.title}前10行未找到指标编号或指标描述列"
                )
            elif header_row == 0:
                errors.append(f"工作表{worksheet.title}未识别到表头")
    finally:
        workbook.close()

    missing_sheets = sorted(set(expected_sheets) - seen_sheets)
    for sheet in missing_sheets:
        errors.append(f"缺少工作表：{CATEGORY_CONFIG[sheet[0]]['sheet']}")

    return indicators, errors, len(seen_sheets)


def load_risk_tree(
    tree_path: str | Path,
) -> tuple[dict[str, Any], dict[str, TreeIndicator], list[str]]:
    path = Path(tree_path)
    if not path.is_file():
        return {}, {}, [f"风险树文件不存在：{path}"]

    try:
        tree = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        return {}, {}, [f"风险树读取失败：{exc}"]

    if not isinstance(tree, dict):
        return {}, {}, ["风险树根节点必须是JSON对象"]

    expected_roots = {config["risk_root"] for config in CATEGORY_CONFIG.values()}
    actual_roots = set(tree)
    errors: list[str] = []
    if actual_roots != expected_roots:
        missing = sorted(expected_roots - actual_roots)
        extra = sorted(actual_roots - expected_roots)
        errors.append(f"风险树根类别不一致，缺少={missing}，多出={extra}")

    root_prefix = {
        config["risk_root"]: prefix for prefix, config in CATEGORY_CONFIG.items()
    }
    indicators: dict[str, TreeIndicator] = {}
    for risk_root, indicator_nodes in tree.items():
        if risk_root not in root_prefix:
            continue
        if not isinstance(indicator_nodes, dict):
            errors.append(f"风险类别{risk_root}的指标集合必须是JSON对象")
            continue

        for indicator_name, leaf in indicator_nodes.items():
            if not isinstance(leaf, dict):
                errors.append(f"指标{indicator_name}的叶节点必须是JSON对象")
                continue
            if "序号" not in leaf or "映射文件" not in leaf:
                errors.append(f"指标{indicator_name}缺少序号或映射文件字段")
                continue

            sequence = _normalize_sequence(leaf["序号"])
            match = SEQUENCE_PATTERN.fullmatch(sequence)
            if match is None:
                errors.append(f"指标{indicator_name}的序号格式错误：{sequence}")
                continue
            if match.group("category") != root_prefix[risk_root]:
                errors.append(
                    f"指标{sequence}所属风险类别错误：当前为{risk_root}"
                )
                continue
            if not isinstance(leaf["映射文件"], dict):
                errors.append(f"指标{sequence}的映射文件必须是JSON对象")
                continue
            if sequence in indicators:
                previous = indicators[sequence]
                errors.append(
                    f"风险树指标编号{sequence}重复：{previous.indicator_name}、{indicator_name}"
                )
                continue

            indicators[sequence] = TreeIndicator(
                sequence=sequence,
                indicator_name=indicator_name,
                risk_root=risk_root,
                leaf=leaf,
            )

    return tree, indicators, errors


def _without_descriptions(value: Any) -> Any:
    if isinstance(value, dict):
        return {
            key: _without_descriptions(item)
            for key, item in value.items()
            if key != "指标描述"
        }
    if isinstance(value, list):
        return [_without_descriptions(item) for item in value]
    return value


def _enrich_tree(
    tree: dict[str, Any],
    workbook_indicators: Mapping[str, WorkbookIndicator],
) -> tuple[dict[str, Any], dict[str, int]]:
    enriched = copy.deepcopy(tree)
    counts = {"added": 0, "updated": 0, "unchanged": 0}

    for indicator_nodes in enriched.values():
        for indicator_name, leaf in indicator_nodes.items():
            sequence = _normalize_sequence(leaf["序号"])
            description = workbook_indicators[sequence].description
            existing = leaf.get("指标描述")
            if existing is None:
                counts["added"] += 1
            elif existing == description:
                counts["unchanged"] += 1
            else:
                counts["updated"] += 1

            rebuilt: dict[str, Any] = {}
            for key, value in leaf.items():
                if key == "指标描述":
                    continue
                rebuilt[key] = value
                if key == "序号":
                    rebuilt["指标描述"] = description
            indicator_nodes[indicator_name] = rebuilt

    if _without_descriptions(tree) != _without_descriptions(enriched):
        raise IndicatorDescriptionSyncError(
            ["补充指标描述后，检测到指标描述以外的风险树内容发生变化"]
        )
    return enriched, counts


def _atomic_write_json(path: Path, value: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = json.dumps(value, ensure_ascii=False, indent=2) + "\n"
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            newline="\n",
            dir=path.parent,
            prefix=f".{path.name}.",
            suffix=".tmp",
            delete=False,
        ) as temporary:
            temporary.write(payload)
            temporary.flush()
            os.fsync(temporary.fileno())
            temporary_path = Path(temporary.name)
        os.replace(temporary_path, path)
    finally:
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()


def sync_indicator_descriptions(
    workbook_path: str | Path,
    tree_path: str | Path = DEFAULT_TREE_PATH,
    *,
    write: bool = False,
) -> dict[str, Any]:
    tree_path = Path(tree_path)
    workbook_indicators, workbook_errors, worksheet_count = load_workbook_indicators(
        workbook_path
    )
    tree, tree_indicators, tree_errors = load_risk_tree(tree_path)
    errors = [*workbook_errors, *tree_errors]

    workbook_sequences = set(workbook_indicators)
    tree_sequences = set(tree_indicators)
    workbook_only = sorted(workbook_sequences - tree_sequences)
    tree_only = sorted(tree_sequences - workbook_sequences)
    if workbook_only or tree_only:
        errors.append(
            f"工作簿与风险树指标编号集合不一致，工作簿独有={workbook_only}，风险树独有={tree_only}"
        )

    for sequence in sorted(workbook_sequences & tree_sequences):
        workbook_item = workbook_indicators[sequence]
        tree_item = tree_indicators[sequence]
        if workbook_item.risk_root != tree_item.risk_root:
            errors.append(
                f"指标{sequence}风险类别不一致：工作簿={workbook_item.risk_root}，"
                f"风险树={tree_item.risk_root}"
            )

    base_report: dict[str, Any] = {
        "mode": "write" if write else "check",
        "workbook": str(Path(workbook_path)),
        "tree": str(tree_path),
        "worksheet_count": worksheet_count,
        "workbook_indicator_count": len(workbook_indicators),
        "tree_indicator_count": len(tree_indicators),
        "workbook_only_sequences": workbook_only,
        "tree_only_sequences": tree_only,
    }
    if errors:
        raise IndicatorDescriptionSyncError(errors, base_report)

    enriched, counts = _enrich_tree(tree, workbook_indicators)
    report = {
        **base_report,
        **counts,
        "errors": [],
        "written": False,
    }
    if write and (counts["added"] or counts["updated"]):
        _atomic_write_json(tree_path, enriched)
        report["written"] = True
    return report


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="从企业合规风险指标工作簿向风险文件映射树同步三级指标描述"
    )
    parser.add_argument("--workbook", required=True, help="企业合规风险指标工作簿路径")
    parser.add_argument(
        "--tree",
        default=str(DEFAULT_TREE_PATH),
        help="风险文件映射树路径，默认使用仓库内risk2file_tree.json",
    )
    parser.add_argument(
        "--write",
        action="store_true",
        help="通过全部校验后原子写入；未指定时仅检查",
    )
    args = parser.parse_args(argv)

    try:
        report = sync_indicator_descriptions(
            args.workbook,
            args.tree,
            write=args.write,
        )
    except IndicatorDescriptionSyncError as exc:
        print(json.dumps(exc.report, ensure_ascii=False, indent=2))
        return 1

    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
