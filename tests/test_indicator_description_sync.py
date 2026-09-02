from __future__ import annotations

import copy
import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from app.infrastructure.kg_to_natural_language.risk_tree.indicator_description_sync import (
    CATEGORY_CONFIG,
    IndicatorDescriptionSyncError,
    sync_indicator_descriptions,
)


def _write_workbook(path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for prefix, config in CATEGORY_CONFIG.items():
        worksheet = workbook.create_sheet(config["sheet"])
        worksheet.append(["说明"])
        worksheet.append(["一级指标", "二级指标", "三级指标", "指标编号", "指标描述"])
        worksheet.append(
            [
                "一级",
                "二级",
                f"{prefix}.1.1.1 指标名称（满分6分）",
                f"{prefix}.1.1.1 ",
                f"第{prefix}类指标的完整描述。\r\n第二行。",
            ]
        )
    workbook.save(path)


def _tree() -> dict[str, object]:
    return {
        config["risk_root"]: {
            f"第{prefix}类简短指标名称": {
                "序号": f"{prefix}.1.1.1",
                "映射文件": {
                    "中国": {
                        f"文件{prefix}": {
                            "原始文本": f"《文件{prefix}》",
                            "已找到": f"文件{prefix}.txt",
                        }
                    }
                },
            }
        }
        for prefix, config in CATEGORY_CONFIG.items()
    }


def _write_tree(path: Path, tree: dict[str, object] | None = None) -> None:
    path.write_text(
        json.dumps(tree or _tree(), ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def _strip_descriptions(value: object) -> object:
    if isinstance(value, dict):
        return {
            key: _strip_descriptions(item)
            for key, item in value.items()
            if key != "指标描述"
        }
    if isinstance(value, list):
        return [_strip_descriptions(item) for item in value]
    return value


def test_check_write_preservation_and_idempotency(tmp_path: Path) -> None:
    workbook_path = tmp_path / "indicators.xlsx"
    tree_path = tmp_path / "tree.json"
    _write_workbook(workbook_path)
    original_tree = _tree()
    _write_tree(tree_path, original_tree)
    original_bytes = tree_path.read_bytes()

    check = sync_indicator_descriptions(workbook_path, tree_path)
    assert check["worksheet_count"] == 6
    assert check["workbook_indicator_count"] == 6
    assert check["added"] == 6
    assert check["written"] is False
    assert tree_path.read_bytes() == original_bytes

    written = sync_indicator_descriptions(workbook_path, tree_path, write=True)
    assert written["added"] == 6
    assert written["updated"] == 0
    assert written["written"] is True

    enriched = json.loads(tree_path.read_text(encoding="utf-8"))
    assert _strip_descriptions(enriched) == original_tree
    for prefix, config in CATEGORY_CONFIG.items():
        leaf = enriched[config["risk_root"]][f"第{prefix}类简短指标名称"]
        assert list(leaf)[:3] == ["序号", "指标描述", "映射文件"]
        assert leaf["指标描述"] == f"第{prefix}类指标的完整描述。\n第二行。"

    repeated = sync_indicator_descriptions(workbook_path, tree_path, write=True)
    assert repeated["added"] == repeated["updated"] == 0
    assert repeated["unchanged"] == 6
    assert repeated["written"] is False


def test_exact_sequence_join_ignores_indicator_name_difference(tmp_path: Path) -> None:
    workbook_path = tmp_path / "indicators.xlsx"
    tree_path = tmp_path / "tree.json"
    _write_workbook(workbook_path)
    _write_tree(tree_path)

    sync_indicator_descriptions(workbook_path, tree_path, write=True)
    enriched = json.loads(tree_path.read_text(encoding="utf-8"))
    first = enriched[CATEGORY_CONFIG["1"]["risk_root"]]["第1类简短指标名称"]
    assert first["指标描述"].startswith("第1类指标的完整描述")


@pytest.mark.parametrize("failure", ["blank", "malformed", "duplicate"])
def test_invalid_workbook_does_not_modify_tree(tmp_path: Path, failure: str) -> None:
    workbook_path = tmp_path / "indicators.xlsx"
    tree_path = tmp_path / "tree.json"
    _write_workbook(workbook_path)
    _write_tree(tree_path)

    workbook = load_workbook(workbook_path)
    worksheet = workbook[CATEGORY_CONFIG["1"]["sheet"]]
    if failure == "blank":
        worksheet.cell(3, 5).value = None
    elif failure == "malformed":
        worksheet.cell(3, 4).value = "1.1"
    else:
        worksheet.append(["一级", "二级", "重复指标", "1.1.1.1", "重复描述"])
    workbook.save(workbook_path)
    before = tree_path.read_bytes()

    with pytest.raises(IndicatorDescriptionSyncError):
        sync_indicator_descriptions(workbook_path, tree_path, write=True)
    assert tree_path.read_bytes() == before


def test_set_and_category_mismatch_do_not_modify_tree(tmp_path: Path) -> None:
    workbook_path = tmp_path / "indicators.xlsx"
    tree_path = tmp_path / "tree.json"
    _write_workbook(workbook_path)

    missing_tree = _tree()
    missing_tree[CATEGORY_CONFIG["1"]["risk_root"]] = {}
    _write_tree(tree_path, missing_tree)
    before = tree_path.read_bytes()
    with pytest.raises(IndicatorDescriptionSyncError, match="指标编号集合不一致"):
        sync_indicator_descriptions(workbook_path, tree_path, write=True)
    assert tree_path.read_bytes() == before

    conflicting_tree = copy.deepcopy(_tree())
    first_root = CATEGORY_CONFIG["1"]["risk_root"]
    second_root = CATEGORY_CONFIG["2"]["risk_root"]
    first_leaf = conflicting_tree[first_root].pop("第1类简短指标名称")
    conflicting_tree[second_root]["错误类别指标"] = first_leaf
    _write_tree(tree_path, conflicting_tree)
    before = tree_path.read_bytes()
    with pytest.raises(IndicatorDescriptionSyncError, match="所属风险类别错误"):
        sync_indicator_descriptions(workbook_path, tree_path, write=True)
    assert tree_path.read_bytes() == before
